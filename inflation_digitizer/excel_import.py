"""Read quarterly central estimates from report-named Excel worksheets."""
import math
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .configuration import country_code
from .forecast_export import REQUIRED, convert_frame, export_csv_files


def clean(value):
    return '' if value is None else ' '.join(str(value).split())


def report_quarter(name):
    text = clean(name)
    match = re.fullmatch(r'(\d{4})[ _-]*(Q[1-4]|March|Mar|June|Jun|September|Sept|Sep|December|Dec)', text, re.I)
    if not match:
        return None
    year, period = match.groups()
    period = period.lower()
    quarter = int(period[1]) if period.startswith('q') else {'mar':1, 'jun':2, 'sep':3, 'dec':4}[period[:3]]
    return f'{year}Q{quarter}'


def target_quarter(year, period):
    year = clean(year)
    period = clean(period).upper()
    period = {'I':'1', 'II':'2', 'III':'3', 'IV':'4'}.get(period, period)
    if not re.fullmatch(r'\d{4}', year) or not re.fullmatch(r'Q?[1-4]', period):
        raise ValueError(f'Invalid target year/quarter: {year!r}, {period!r}')
    return f'{year}Q{period[-1]}'


def read_excel_tables(path, country):
    """Return detailed source rows and names of skipped non-report worksheets."""
    path = Path(path)
    country = country_code(country)
    rows, skipped = [], []
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        for sheet in workbook:
            report = report_quarter(sheet.title)
            if report is None:
                # Never silently skip an apparent dated report with an unsupported name.
                if re.search(r'\d{4}', sheet.title):
                    raise ValueError(f'{sheet.title}: use a report name such as 2020 Mar or 2020Q1')
                skipped.append(sheet.title)
                continue
            iterator = sheet.iter_rows()
            header = next(iterator, ())
            names = [clean(cell.value).lower() for cell in header]
            for required in ('period', 'central'):
                if names.count(required) != 1:
                    raise ValueError(f'{sheet.title}: expected one {required!r} column')
            if names.count('year') > 1:
                raise ValueError(f'{sheet.title}: duplicate Year columns')
            pc, vc = names.index('period'), names.index('central')
            yc = names.index('year') if 'year' in names else None
            count = 0
            for cells in iterator:
                values = [cell.value for cell in cells]
                if all(value is None or clean(value) == '' for value in values):
                    continue
                location = f'{sheet.title}!{cells[vc].coordinate}'
                try:
                    if yc is None:
                        parts = clean(values[pc]).split()
                        if len(parts) != 2:
                            raise ValueError('Period must contain year and quarter, e.g. 2015 2')
                        target = target_quarter(*parts)
                    else:
                        target = target_quarter(values[yc], values[pc])
                    raw = values[vc]
                    if cells[vc].data_type == 'f' or not isinstance(raw, (int, float, str)) or isinstance(raw, bool):
                        raise ValueError('central must be a numeric value, not a formula or blank')
                    value = float(clean(raw))
                    if not math.isfinite(value):
                        raise ValueError('central must be finite')
                    if '%' in cells[vc].number_format:
                        raise ValueError('central must be in percentage points (e.g. 8.1), not Excel percent format')
                except (ValueError, TypeError) as exc:
                    raise ValueError(f'{location}: {exc}') from exc
                rows.append(dict(country=country, vintage=report, quarter=target,
                                 inflation_pct_yoy=value, series='reported_central',
                                 source_type='excel_table', source_file=str(path.resolve()),
                                 source_sheet=sheet.title, source_cell=cells[vc].coordinate))
                count += 1
            if count == 0:
                raise ValueError(f'{sheet.title}: no estimates found')
    finally:
        workbook.close()
    if not rows:
        raise ValueError('No report sheets found; use sheet names such as 2020 Mar or 2020Q1')
    frame = pd.DataFrame(rows)
    convert_frame(frame)  # Validate dates and duplicate report/target keys.
    return frame, skipped


def import_excel_tables(path, country, output, variable='CPI'):
    """Persist one workbook as a source CSV used by all subsequent exports."""
    path, output = Path(path), Path(output)
    frame, skipped = read_excel_tables(path, country)
    if not re.fullmatch(r'[A-Za-z0-9_]+', variable):
        raise ValueError('Variable must contain only letters, digits, and underscores')
    tables = output/'tables'
    destination = tables/f'{country_code(country)}_{variable}_excel.csv'
    # Refuse to replace a different workbook accidentally. Reimporting the same
    # workbook replaces its full snapshot, so removed rows do not linger.
    if destination.exists():
        previous = pd.read_csv(destination)
        if ('source_file' not in previous or 'source_type' not in previous or
                not previous.source_file.eq(str(path.resolve())).all() or
                not previous.source_type.eq('excel_table').all()):
            raise ValueError(f'{destination}: already belongs to another source')
    others = [pd.read_csv(p) for p in sorted(tables.glob('*.csv')) if p != destination]
    convert_frame(pd.concat([*others, frame], ignore_index=True), variable)
    tables.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix('.csv.tmp')
    frame.to_csv(temporary, index=False)
    temporary.replace(destination)
    combined = export_csv_files(tables, output/'forecasts.csv', variable, output/'formatted_tables')
    return frame, combined, skipped
