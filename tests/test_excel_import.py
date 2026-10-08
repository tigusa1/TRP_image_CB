import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook

from inflation_digitizer.excel_import import read_excel_tables, import_excel_tables
from inflation_digitizer.forecast_export import export_csv_files


class ExcelImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.path = self.root/'BRA CPI.xlsx'

    def tearDown(self):
        self.temp.cleanup()

    def workbook(self, rows, title='2020 Mar'):
        book = Workbook()
        sheet = book.active
        sheet.title = title
        for row in rows:
            sheet.append(row)
        return book

    def read(self, book):
        with patch('inflation_digitizer.excel_import.load_workbook', return_value=book):
            return read_excel_tables(self.path, 'BRA')

    def test_layouts_whitespace_roman_quarters_and_horizons(self):
        book = self.workbook([['Year','Period','central'],[2019,'IV',4.3],
                              [2020,'I',0],[2020,'II','3.0\u00a0']])
        sheet = book.create_sheet('2020 Jun')
        sheet.append(['Period',-.5,'central'])
        sheet.append(['2020 2\u00a0',99,'2.1\u00a0'])
        book.create_sheet('wraprows').append(['not a report'])
        frame, skipped = self.read(book)
        from inflation_digitizer.forecast_export import convert_frame
        result = convert_frame(frame)
        self.assertEqual(result.h.tolist(),[-1,0,1,0])
        self.assertEqual(result.forecast.tolist(),[4.3,0,3,2.1])
        self.assertEqual(skipped,['wraprows'])
        self.assertEqual(frame.source_cell.tolist(),['C2','C3','C4','C2'])

    def test_invalid_data_does_not_become_zero(self):
        for value in (None, 'bad', '=1+1', float('inf')):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError,'C2'):
                self.read(self.workbook([['Year','Period','central'],[2020,'I',value]]))

    def test_duplicate_targets_and_unsupported_report_names_fail(self):
        with self.assertRaisesRegex(ValueError,'Duplicate'):
            self.read(self.workbook([['Year','Period','central'],[2020,'I',1],[2020,'I',2]]))
        with self.assertRaisesRegex(ValueError,'report name'):
            self.read(self.workbook([['Year','Period','central']],title='2020 April'))

    def test_reimport_and_normal_refresh_preserve_sources(self):
        book = self.workbook([['Year','Period','central'],[2020,'I',3.4],[2020,'II',3]])
        frame, _ = self.read(book)
        out = self.root/'output'
        tables = out/'tables'
        tables.mkdir(parents=True)
        chart = pd.DataFrame(dict(country=['BRA'], vintage=['2019Q4'], quarter=['2020Q1'], inflation_pct_yoy=[3.7]))
        chart.to_csv(tables/'chart.csv',index=False)
        chart_bytes = (tables/'chart.csv').read_bytes()
        with patch('inflation_digitizer.excel_import.read_excel_tables',return_value=(frame,[])):
            _, combined, _ = import_excel_tables(self.path,'BRA',out)
            self.assertEqual(len(combined),3)
            import_excel_tables(self.path,'BRA',out)
        changed = frame.iloc[:1].copy()
        changed['inflation_pct_yoy'] = 3.5
        with patch('inflation_digitizer.excel_import.read_excel_tables',return_value=(changed,[])):
            import_excel_tables(self.path,'BRA',out)
        result = export_csv_files(tables,out/'forecasts.csv')
        self.assertEqual(len(result),2)
        self.assertEqual(result.forecast.tolist(),[3.7,3.5])
        self.assertEqual((tables/'chart.csv').read_bytes(),chart_bytes)

    def test_conflict_rejected_before_existing_files_change(self):
        frame, _ = self.read(self.workbook([['Year','Period','central'],[2020,'I',3.4]]))
        out = self.root/'output'
        (out/'tables').mkdir(parents=True)
        frame.to_csv(out/'tables'/'chart.csv',index=False)
        (out/'forecasts.csv').write_text('previous output')
        with patch('inflation_digitizer.excel_import.read_excel_tables',return_value=(frame,[])):
            with self.assertRaisesRegex(ValueError,'Duplicate'):
                import_excel_tables(self.path,'BRA',out)
        self.assertFalse((out/'tables'/'BRA_CPI_excel.csv').exists())
        self.assertEqual((out/'forecasts.csv').read_text(),'previous output')


if __name__ == '__main__':
    unittest.main()
